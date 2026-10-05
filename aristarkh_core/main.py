import asyncio
import os
import json
import aiohttp 
import nest_asyncio
import base64
import tempfile
import pypdf
import docx
import openpyxl
from pptx import Presentation
import html 
import time  # [NEW] Внутренняя шкала времени Аристарха

from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import Command, CommandObject, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.fsm.storage.redis import RedisStorage 
from redis.asyncio import Redis 
from aiogram.types import (
    ReplyKeyboardMarkup, KeyboardButton, 
    InlineKeyboardMarkup, InlineKeyboardButton, 
    FSInputFile, LabeledPrice, PreCheckoutQuery, LinkPreviewOptions
)
from aiogram.exceptions import TelegramBadRequest 
from aiogram.client.session.aiohttp import AiohttpSession

from aristarkh_core.config import config
import logging
logger = logging.getLogger(__name__)
import i18n
from aristarkh_core.prompts import Prompts
from aristarkh_core.database import BRIEF_FIELDS, DatabaseService
from aristarkh_core.reply_guard import polish_reply
from aristarkh_core.rag_chroma import RAGService
from aristarkh_core.gemini import LLMService
from aristarkh_core.assistant import AssistantService 
from semantic_memory import SemanticMemory
from aristarkh_core.evolution_engine import PersonaEvolutionEngine
from aristarkh_core.context_simulator import ContextSimulator  # [NEW] Симулятор среды обитания
from news_module.post_builder import BOT_CTA_ENABLED, extract_url, generate_post, generate_razbor
from aristarkh_core.humanize import humanize_punctuation
import uuid

# --- INIT ---
try:
    asyncio.set_event_loop_policy(asyncio.DefaultEventLoopPolicy())
except Exception as e:
    logger.warning(f"Не удалось применить DefaultEventLoopPolicy: {e}")

nest_asyncio.apply()

# [FIX] Увеличиваем таймаут и добавляем поддержку прокси из .env
proxy_url = os.getenv("PROXY_URL")
if proxy_url:
    logger.info(f"🌐 Включен прокси-сервер для обхода блокировок Telegram")
    bot_session = AiohttpSession(timeout=120, proxy=proxy_url)
else:
    bot_session = AiohttpSession(timeout=120)

bot = Bot(token=config.BOT_TOKEN, session=bot_session)

if config.USE_REDIS:
    logger.info("🔌 Подключаем RedisStorage для стейтов...")
    redis_client = Redis.from_url(config.REDIS_URL, decode_responses=True)
    storage = RedisStorage(redis_client)
else:
    logger.info("🧠 Используется MemoryStorage (Режим Colab). Для прода включите USE_REDIS.")
    storage = MemoryStorage()

dp = Dispatcher(storage=storage)
db_service = DatabaseService(config.DB_URL)
rag_service = RAGService(config.GOOGLE_API_KEY)
llm_service = LLMService(config.GOOGLE_API_KEY)
assistant_service = AssistantService(config.GOOGLE_API_KEY) 
semantic_memory = SemanticMemory()
evolution_engine = PersonaEvolutionEngine(config.GOOGLE_API_KEY)

# [КРИТИЧЕСКИЙ ФИКС] Абсолютный путь
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
AGENDA_PATH = os.path.join(BASE_DIR, "aristarkh_agenda.json")

# Канал для авторских постов Аристарха (задаётся только через .env)
CHANNEL_ID = os.getenv("TG_CHANNEL_ID", "")

# --- STATES ---
class BotStates(StatesGroup):
    chatting = State()
    generating = State()  # [NEW] Режим блокировки — Аристарх думает
    waiting_for_long_input = State()
    waiting_for_promo = State()


class BriefStates(StatesGroup):
    """Бриф проекта: по одному состоянию на вопрос, порядок как в BRIEF_FIELDS."""
    niche = State()
    audience = State()
    goal = State()
    channel = State()
    constraints = State()


BRIEF_STEPS = [getattr(BriefStates, f) for f in BRIEF_FIELDS]

# --- [NEW] ГЛОБАЛЬНЫЕ БУФЕРЫ ДЛЯ DEBOUNCE (Анти-пулемет) ---
debounce_tasks = {}
user_message_buffers = {}
user_image_buffers = {}

# --- KEYBOARDS ---
def main_kb(lang: str) -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(keyboard=[
        [
            KeyboardButton(text=i18n.get_text("btn_brief", lang)),
            KeyboardButton(text=i18n.get_text("btn_settings", lang)),
        ],
        [
            KeyboardButton(text=i18n.get_text("btn_balance", lang)),
            KeyboardButton(text=i18n.get_text("btn_rest", lang)),
        ],
    ], resize_keyboard=True)


def cancel_kb(lang: str) -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(keyboard=[[KeyboardButton(text=i18n.get_text("btn_main_menu", lang))]], resize_keyboard=True)


def long_input_kb(lang: str) -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(keyboard=[[
        KeyboardButton(text=i18n.get_text("btn_done", lang)),
        KeyboardButton(text=i18n.get_text("btn_cancel", lang)),
    ]], resize_keyboard=True)


def settings_kb(t: float, lang: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"{'✅ ' if t==0.2 else ''}{i18n.get_text('temp_dry', lang)}", callback_data="temp_0.2")],
        [InlineKeyboardButton(text=f"{'✅ ' if t==0.7 else ''}{i18n.get_text('temp_norm', lang)}", callback_data="temp_0.7")],
        [InlineKeyboardButton(text=f"{'✅ ' if t==1.3 else ''}{i18n.get_text('temp_boom', lang)}", callback_data="temp_1.3")],
        [InlineKeyboardButton(text=i18n.get_text("btn_reset_context", lang), callback_data="reset_context")]
    ])


MENU_BUTTONS = (i18n.variants("btn_settings") | i18n.variants("btn_balance") | i18n.variants("btn_rest")
                | i18n.variants("btn_brief"))
DONE_WORDS = {"все", "всё", "done", "готово"} | {v.lower() for v in i18n.variants("btn_done")}

# --- UTILS: TEXT EXTRACTOR ---
async def extract_text_from_file(file_path: str, ext: str) -> str:
    try:
        if ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f: return f.read()
        elif ext == '.pdf':
            reader = pypdf.PdfReader(file_path)
            return "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
        elif ext == '.docx':
            doc = docx.Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs])
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_texts = [str(cell) for cell in row if cell is not None]
                    if row_texts: text.append(" ".join(row_texts))
            return "\n".join(text)
        elif ext == '.pptx':
            prs = Presentation(file_path)
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return "\n".join(text)
    except Exception as e:
        logger.error(f"File parsing error: {e}")
        return f"⚠️ Ошибка чтения файла: {e}"
    return ""

# --- UTILS: FILE SENDER (SAFE MODE) ---
async def send_smart_response(message: types.Message, text: str):
    # [NEW] Аппаратная зачистка случайных звёздочек для будущего TTS
    text = text.replace('*', '')
    
    if len(text) > 4000:
        filename = f"response_{message.from_user.id}.txt"
        clean_text = text.replace("__", "")
        
        with open(filename, "w", encoding="utf-8") as f:
            f.write(clean_text)
        
        try:
            await message.answer_document(
                FSInputFile(filename), 
                caption=i18n.get_text("file_too_big", i18n.user_lang(message.from_user))
            )
        finally:
            if os.path.exists(filename):
                os.remove(filename)
    else:
        try:
            await message.answer(text, parse_mode="Markdown")
        except TelegramBadRequest:
            logger.warning(f"Markdown failed for user {message.from_user.id}, sending plain text.")
            await message.answer(text, parse_mode=None)
        except Exception as e:
            logger.error(f"Send error: {e}")
            await message.answer(i18n.get_text("error_sending", i18n.user_lang(message.from_user)))

# --- [NEW] DEBOUNCE TRIGGER TASK ---
async def trigger_processing(message: types.Message, state: FSMContext, user_id: int):
    """Задача таймера. Ждёт тишины 2.5 секунды, затем отправляет склеенный текст в LLM."""
    try:
        await asyncio.sleep(2.5)
    except asyncio.CancelledError:
        return
        
    final_text = user_message_buffers.pop(user_id, "")
    final_img = user_image_buffers.pop(user_id, None)
    
    if not final_text and not final_img:
        return

    await state.set_state(BotStates.generating)
    try:
        await process_ai_response(message, state, final_text, final_img)
    finally:
        current_state = await state.get_state()
        if current_state == BotStates.generating.state:
            await state.set_state(BotStates.chatting)

# --- HANDLERS: START ---
@dp.message(Command("start"))
async def cmd_start(message: types.Message, state: FSMContext):
    await state.clear()
    user = await db_service.get_user(message.from_user.id)
    await db_service.log_event(message.from_user.id, "START_COMMAND", "New user" if not user else "Returning user")

    if not user:
        user = await db_service.create_user(message.from_user.id, message.from_user.username, message.from_user.full_name)
        safe_name = html.escape(message.from_user.full_name or "Без имени")
        safe_username = html.escape(message.from_user.username or "без_юзернейма")
        try:
            await bot.send_message(
                config.ADMIN_ID, 
                i18n.get_text("new_user_alert", user_id=message.from_user.id, username=safe_username, name=safe_name),
                parse_mode="HTML"
            )
        except Exception as e:
            logger.error(f"Не удалось отправить уведомление админу: {e}")
    
    lang = i18n.user_lang(message.from_user)
    welcome_text = i18n.get_text("start_message", lang)
    await state.set_state(BotStates.chatting)
    await state.update_data(interaction_count=0)
    await message.answer(welcome_text, reply_markup=main_kb(lang), parse_mode="Markdown")
    if not await db_service.get_brief(message.from_user.id):
        await message.answer(
            i18n.get_text("brief_offer", lang),
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text=i18n.get_text("btn_brief_start", lang), callback_data="brief_start")
            ]]),
        )

# --- БРИФ ПРОЕКТА: 5 вопросов, ответы подмешиваются в контекст каждого запроса ---
# Бриф не тратит запросы: это обычные сообщения без вызова модели.
BRIEF_SKIP_WORDS = {"-", "—", "нет", "пропустить", "skip", "no"}
BRIEF_LABELS = {
    "niche": "Ниша и чем занимается",
    "audience": "Аудитория",
    "goal": "Цель на 1-3 месяца",
    "channel": "Канал или площадка",
    "constraints": "Ограничения и пожелания",
}


def format_brief(brief: dict | None) -> str:
    lines = [f"- {BRIEF_LABELS[f]}: {brief[f]}" for f in BRIEF_FIELDS if brief and brief.get(f)]
    return "[БРИФ ПРОЕКТА — опирайся на него в первую очередь]\n" + "\n".join(lines) if lines else ""


async def start_brief(message: types.Message, state: FSMContext, lang: str):
    await db_service.log_event(message.chat.id, "BRIEF_STARTED")
    await state.set_state(BRIEF_STEPS[0])
    await state.update_data(brief_answers={})
    await message.answer(i18n.get_text("brief_intro", lang), reply_markup=cancel_kb(lang))
    await message.answer(i18n.get_text(f"brief_q_{BRIEF_FIELDS[0]}", lang))


@dp.message(Command("brief"))
async def cmd_brief(message: types.Message, state: FSMContext):
    await start_brief(message, state, i18n.user_lang(message.from_user))


@dp.message(F.text.in_(i18n.variants("btn_brief")))
async def brief_button(message: types.Message, state: FSMContext):
    await start_brief(message, state, i18n.user_lang(message.from_user))


@dp.callback_query(F.data == "brief_start")
async def brief_start_callback(callback: types.CallbackQuery, state: FSMContext):
    await callback.answer()
    await start_brief(callback.message, state, i18n.user_lang(callback.from_user))


@dp.message(StateFilter(*BRIEF_STEPS), F.text)
async def brief_answer(message: types.Message, state: FSMContext):
    lang = i18n.user_lang(message.from_user)
    if message.text in i18n.variants("btn_main_menu"):
        await state.set_state(BotStates.chatting)
        await message.answer(i18n.get_text("brief_cancelled", lang), reply_markup=main_kb(lang))
        return
    step = [s.state for s in BRIEF_STEPS].index(await state.get_state())
    answers = (await state.get_data()).get("brief_answers", {})
    answers[BRIEF_FIELDS[step]] = "" if message.text.strip().lower() in BRIEF_SKIP_WORDS else message.text
    if step + 1 < len(BRIEF_STEPS):
        await state.update_data(brief_answers=answers)
        await state.set_state(BRIEF_STEPS[step + 1])
        await message.answer(i18n.get_text(f"brief_q_{BRIEF_FIELDS[step + 1]}", lang))
        return
    await db_service.save_brief(message.from_user.id, answers)
    await db_service.log_event(message.from_user.id, "BRIEF_SAVED")
    await state.update_data(brief_answers={})
    await state.set_state(BotStates.chatting)
    await message.answer(i18n.get_text("brief_saved", lang), reply_markup=main_kb(lang))

# --- РУБРИКА «РАЗБОР ПОДПИСЧИКА»: только администратор ---
# /razbor <вопрос> → Аристарх пишет разбор → превью с кнопками → публикация в канал.
# Черновики живут в памяти процесса: после перезапуска бота превью нужно запросить заново.
razbor_drafts: dict[str, dict] = {}


def razbor_kb(draft_id: str, lang: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text=i18n.get_text("btn_razbor_publish", lang), callback_data=f"rz_pub:{draft_id}"),
        InlineKeyboardButton(text=i18n.get_text("btn_razbor_redo", lang), callback_data=f"rz_redo:{draft_id}"),
        InlineKeyboardButton(text=i18n.get_text("btn_razbor_cancel", lang), callback_data=f"rz_del:{draft_id}"),
    ]])


async def make_razbor(question: str) -> str:
    # Ссылка на бота только при CHANNEL_BOT_CTA; иначе разбор зовёт писать вопросы в комментарии
    cta_handle = f"@{(await bot.get_me()).username}" if BOT_CTA_ENABLED else None
    core_beliefs = evolution_engine.get_current_beliefs()
    post_html, _ = await generate_razbor(llm_service, rag_service, core_beliefs, question, cta_handle)
    return post_html


@dp.message(Command("razbor"), F.from_user.id == config.ADMIN_ID)
async def cmd_razbor(message: types.Message, command: CommandObject):
    lang = i18n.user_lang(message.from_user)
    question = (command.args or "").strip()
    if not question:
        await message.answer(i18n.get_text("razbor_usage", lang))
        return
    if not CHANNEL_ID:
        await message.answer(i18n.get_text("razbor_no_channel", lang))
        return
    wait_msg = await message.answer(i18n.get_text("razbor_wait", lang))
    try:
        post_html = await make_razbor(question)
    except Exception as e:
        logger.error(f"Razbor error: {e}")
        await wait_msg.edit_text(i18n.get_text("post_error", lang, error=e))
        return
    draft_id = uuid.uuid4().hex[:8]
    razbor_drafts[draft_id] = {"question": question, "html": post_html}
    await wait_msg.delete()
    await message.answer(
        post_html, parse_mode="HTML",
        link_preview_options=LinkPreviewOptions(is_disabled=True),
        reply_markup=razbor_kb(draft_id, lang),
    )


@dp.callback_query(F.data.startswith("rz_"), F.from_user.id == config.ADMIN_ID)
async def razbor_action(callback: types.CallbackQuery):
    lang = i18n.user_lang(callback.from_user)
    action, _, draft_id = callback.data.partition(":")
    if draft_id not in razbor_drafts:
        await callback.answer(i18n.get_text("razbor_expired", lang), show_alert=True)
        return

    if action == "rz_pub":
        draft = razbor_drafts.pop(draft_id)  # сразу забираем: двойное нажатие не опубликует дважды
        try:
            await bot.send_message(
                CHANNEL_ID, draft["html"], parse_mode="HTML",
                link_preview_options=LinkPreviewOptions(is_disabled=True),
            )
        except Exception as e:
            razbor_drafts[draft_id] = draft
            await callback.answer(i18n.get_text("post_error", lang, error=e), show_alert=True)
            return
        await db_service.log_event(callback.from_user.id, "RAZBOR_PUBLISHED", f"channel:{CHANNEL_ID}")
        await callback.message.edit_reply_markup(reply_markup=None)
        await callback.message.reply(i18n.get_text("razbor_published", lang, channel=CHANNEL_ID))
        await callback.answer()
    elif action == "rz_redo":
        await callback.answer(i18n.get_text("razbor_wait", lang))
        try:
            post_html = await make_razbor(razbor_drafts[draft_id]["question"])
        except Exception as e:
            logger.error(f"Razbor redo error: {e}")
            await callback.message.reply(i18n.get_text("post_error", lang, error=e))
            return
        razbor_drafts[draft_id]["html"] = post_html
        await callback.message.edit_text(
            post_html, parse_mode="HTML",
            link_preview_options=LinkPreviewOptions(is_disabled=True),
            reply_markup=razbor_kb(draft_id, lang),
        )
    else:
        razbor_drafts.pop(draft_id, None)
        await callback.message.edit_reply_markup(reply_markup=None)
        await callback.message.reply(i18n.get_text("razbor_cancelled", lang))
        await callback.answer()

# --- HANDLER: ПЕРЕКУР (AFFECTIVE MEMORY RESET — ИСТОРИЯ НЕ СТИРАЕТСЯ) ---
@dp.message(F.text.in_(i18n.variants("btn_rest")))
async def cmd_rest(message: types.Message, state: FSMContext):
    await db_service.log_event(message.from_user.id, "REST_BUTTON_CLICKED")
    await db_service.reset_fatigue(message.from_user.id)
    
    logger.info(f"🚬 [СБРОС] Юзер {message.from_user.id} нажал 'Перекур'. Усталость обнулена, история диалога сохранена.")
    
    await message.answer(i18n.get_text("rest_btn_clicked", i18n.user_lang(message.from_user)), parse_mode="Markdown")

# --- HANDLERS: BALANCE & STARS ---
@dp.message(F.text.in_(i18n.variants("btn_balance")))
async def show_balance(message: types.Message):
    await db_service.log_event(message.from_user.id, "CHECK_BALANCE")
    lang = i18n.user_lang(message.from_user)
    
    user = await db_service.get_user(message.from_user.id)
    if not user:
        user = await db_service.create_user(message.from_user.id, message.from_user.username, message.from_user.full_name)
        
    is_unlim = await db_service.has_unlimited(message.from_user.id)
    
    expired = await db_service.reset_balance_if_expired(user.id)
    if expired:
        user.credits = 0
        user.expires_at = None

    if is_unlim:
        unlim_str = i18n.get_text("balance_unlim_forever", lang) if not user.unlimited_until else user.unlimited_until.strftime('%Y-%m-%d %H:%M')
        text = i18n.get_text("balance_info_unlim", lang, unlim_str=unlim_str, credits=user.credits)
    else:
        expiry_info = i18n.get_text("subscription_until", lang, date=user.expires_at.strftime('%Y-%m-%d %H:%M')) if user.expires_at else ""
        text = i18n.get_text("balance_info", lang, credits=user.credits, expiry_info=expiry_info)
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=i18n.get_text("pkg_10", lang), callback_data="buy_10_0_1000")],
        [InlineKeyboardButton(text=i18n.get_text("pkg_7d", lang), callback_data="buy_25_7_2000")],
        [InlineKeyboardButton(text=i18n.get_text("pkg_30d", lang), callback_data="buy_100_30_5000")],
        [InlineKeyboardButton(text=i18n.get_text("btn_promo", lang), callback_data="enter_promo")]
    ])
    await message.answer(text, reply_markup=kb, parse_mode="Markdown")

@dp.callback_query(F.data == "enter_promo")
async def enter_promo_handler(callback: types.CallbackQuery, state: FSMContext):
    await db_service.log_event(callback.from_user.id, "CLICK_PROMO_BTN")
    await state.set_state(BotStates.waiting_for_promo)
    lang = i18n.user_lang(callback.from_user)
    await callback.message.answer(i18n.get_text("promo_enter", lang), reply_markup=cancel_kb(lang))
    await callback.answer()

@dp.message(BotStates.waiting_for_promo, F.text)
async def process_promo(message: types.Message, state: FSMContext):
    lang = i18n.user_lang(message.from_user)
    if message.text in i18n.variants("btn_main_menu"):
        await db_service.log_event(message.from_user.id, "CANCEL_PROMO_INPUT")
        await state.set_state(BotStates.chatting)
        await message.answer(i18n.get_text("menu_return", lang), reply_markup=main_kb(lang))
        return
        
    if config.PROMO_CODE and message.text.strip().lower() == config.PROMO_CODE.lower():
        await db_service.log_event(message.from_user.id, "ENTER_PROMO_SUCCESS", message.text)
        success = await db_service.activate_promo(message.from_user.id, config.PROMO_HOURS)
        await state.set_state(BotStates.chatting)
        if success:
            await message.answer(i18n.get_text("promo_success", lang), parse_mode="Markdown", reply_markup=main_kb(lang))
        else:
            await message.answer(i18n.get_text("promo_error", lang), reply_markup=main_kb(lang))
    else:
        await db_service.log_event(message.from_user.id, "ENTER_PROMO_FAIL", message.text)
        await state.set_state(BotStates.chatting)
        await message.answer(i18n.get_text("promo_invalid", lang), reply_markup=main_kb(lang))

@dp.callback_query(F.data.startswith("buy_"))
async def buy_stars(callback: types.CallbackQuery):
    _, credits_amount, days, stars = callback.data.split("_")
    await db_service.log_event(callback.from_user.id, "CLICK_BUY_STARS", f"credits:{credits_amount}, days:{days}, stars:{stars}")

    lang = i18n.user_lang(callback.from_user)
    prices = [LabeledPrice(label=i18n.get_text("invoice_label", lang, credits=credits_amount), amount=int(stars))]
    
    await bot.send_invoice(
        chat_id=callback.from_user.id,
        title=i18n.get_text("invoice_title", lang),
        description=i18n.get_text(
            "invoice_desc", lang, credits=credits_amount,
            period=i18n.get_text("invoice_period_days", lang, days=days) if int(days) > 0 else i18n.get_text("invoice_period_forever", lang),
        ),
        payload=callback.data, 
        provider_token="", 
        currency="XTR",
        prices=prices,
        start_parameter="topup"
    )
    await callback.answer()

@dp.pre_checkout_query()
async def pre_checkout_handler(pre_checkout_query: PreCheckoutQuery):
    await bot.answer_pre_checkout_query(pre_checkout_query.id, ok=True)

@dp.message(F.successful_payment)
async def successful_payment(message: types.Message):
    pmt = message.successful_payment
    _, credits_to_add, add_days, amount_stars = pmt.invoice_payload.split("_")
    credits_to_add, add_days, amount_stars = int(credits_to_add), int(add_days), int(amount_stars)
    charge_id = pmt.telegram_payment_charge_id
    
    if await db_service.check_tx_exists(charge_id):
        return

    await db_service.log_event(message.from_user.id, "PAYMENT_SUCCESS", f"stars:{amount_stars}, credits:{credits_to_add}, days:{add_days}")
    new_bal = await db_service.update_balance(message.from_user.id, credits_to_add, add_days)
    await db_service.add_stars_transaction(message.from_user.id, amount_stars, credits_to_add, charge_id)
    await message.answer(i18n.get_text("payment_success", i18n.user_lang(message.from_user), new_bal=new_bal))

# --- HANDLERS: SETTINGS & RESET CONTEXT ---
@dp.message(F.text.in_(i18n.variants("btn_settings")))
async def settings(message: types.Message):
    await db_service.log_event(message.from_user.id, "OPEN_SETTINGS")
    
    user = await db_service.get_user(message.from_user.id)
    if not user:
        user = await db_service.create_user(message.from_user.id, message.from_user.username, message.from_user.full_name)
        
    lang = i18n.user_lang(message.from_user)
    await message.answer(i18n.get_text("settings_title", lang), reply_markup=settings_kb(user.temperature, lang))

@dp.callback_query(F.data == "reset_context")
async def reset_context_handler(callback: types.CallbackQuery, state: FSMContext):
    await db_service.log_event(callback.from_user.id, "RESET_CONTEXT")
    await state.update_data(history=[], interaction_count=0)
    await callback.answer(i18n.get_text("memory_cleared", i18n.user_lang(callback.from_user)), show_alert=True)

@dp.callback_query(F.data.startswith("temp_"))
async def set_temp(callback: types.CallbackQuery):
    t = float(callback.data.split("_")[1])
    user_id = callback.from_user.id
    
    await db_service.log_event(user_id, "CHANGE_TEMP", str(t))
    await db_service.update_temp(user_id, t)
    
    lang = i18n.user_lang(callback.from_user)
    try:
        await callback.message.edit_text(i18n.get_text("settings_title", lang), reply_markup=settings_kb(t, lang))
    except TelegramBadRequest:
        pass
        
    await callback.answer(i18n.get_text("temp_applied", lang))

# --- [NEW] ЗАЩИТА ОТ ПУЛЕМЕТА ВО ВРЕМЯ ГЕНЕРАЦИИ ---
@dp.message(BotStates.generating)
async def generating_handler(message: types.Message):
    await db_service.log_event(message.from_user.id, "SPAM_BLOCKED")
    await message.answer(i18n.get_text("wait_analyzing", i18n.user_lang(message.from_user)))

# --- HANDLERS: CHAT LOGIC (TEXT, PHOTO, DOCS) ---
@dp.message(BotStates.chatting, F.text | F.photo | F.document)
async def chat_handler(message: types.Message, state: FSMContext):
    lang = i18n.user_lang(message.from_user)
    if message.text in i18n.variants("btn_main_menu"):
        await db_service.log_event(message.from_user.id, "RETURN_TO_MAIN_MENU")
        await cmd_start(message, state)
        return
    if message.text in MENU_BUTTONS: return

    # ╔══════════════════════════════════════════════════════════════╗
    # ║  ПУБЛИКАЦИЯ В КАНАЛ ВРУЧНУЮ: только администратор (ADMIN_ID)
    # ╚══════════════════════════════════════════════════════════════╝
    if (message.text and message.text.startswith("[ПУБЛИКАЦИЯ В КАНАЛ]")
            and message.from_user.id == config.ADMIN_ID and CHANNEL_ID):
        await db_service.log_event(message.from_user.id, "CHANNEL_POST_REQUEST")
        wait_msg = await message.answer(i18n.get_text("wait_writing_post", lang))
        
        news_text = message.text.replace("[ПУБЛИКАЦИЯ В КАНАЛ]", "").strip()
        url = extract_url(news_text)
        news = {
            "content": news_text,
            "url": url,
            "original_text": news_text.replace(url, "").strip() if url else news_text,
        }

        try:
            # Тот же конвейер, что и у publisher_cron: формат, тон, анти-клише, RAG, шапка
            core_beliefs = evolution_engine.get_current_beliefs()
            channel_post, post_meta = await generate_post(llm_service, rag_service, core_beliefs, news, [])

            await bot.send_message(
                CHANNEL_ID,
                channel_post,
                parse_mode="HTML",
                link_preview_options=LinkPreviewOptions(is_disabled=True),
            )

            await wait_msg.edit_text(i18n.get_text("post_published", lang, channel=CHANNEL_ID, fmt=post_meta['format']))
            await db_service.log_event(message.from_user.id, "CHANNEL_POST_PUBLISHED", f"channel:{CHANNEL_ID}, chars:{len(channel_post)}, format:{post_meta['format']}")
        except Exception as e:
            logger.error(f"Channel post error: {e}")
            await db_service.log_event(message.from_user.id, "CHANNEL_POST_ERROR", str(e))
            await wait_msg.edit_text(i18n.get_text("post_error", lang, error=e))
        return

    event_type = "INCOMING_REQUEST_TEXT"
    if message.photo: event_type = "INCOMING_REQUEST_PHOTO"
    if message.document: event_type = "INCOMING_REQUEST_DOC"
    await db_service.log_event(message.from_user.id, event_type)

    user_query = message.text or message.caption or "Проанализируй этот файл."
    image_data = None
    file_context = ""

    if message.photo:
        wait_msg = await message.answer(i18n.get_text("wait_looking_pic", lang))
        photo = message.photo[-1]
        file_info = await bot.get_file(photo.file_id)
        downloaded_file = await bot.download_file(file_info.file_path)
        image_bytes = downloaded_file.read()
        image_data = {
            "mime_type": "image/jpeg",
            "data": base64.b64encode(image_bytes).decode("utf-8")
        }
        await wait_msg.delete()

    elif message.document:
        doc = message.document
        ext = os.path.splitext(doc.file_name)[1].lower()
        supported_exts = ['.txt', '.pdf', '.docx', '.xlsx', '.pptx']
        
        await db_service.log_event(message.from_user.id, "FILE_UPLOAD", ext)

        if ext not in supported_exts:
            await message.answer(i18n.get_text("unsupported_format", lang, ext=ext))
            return

        wait_msg = await message.answer(i18n.get_text("wait_reading_doc", lang, ext=ext))
        file_info = await bot.get_file(doc.file_id)
        
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
            await bot.download_file(file_info.file_path, temp_file.name)
            temp_path = temp_file.name

        try:
            extracted_text = await extract_text_from_file(temp_path, ext)
            if extracted_text.strip():
                file_context = f"\n\n[СОДЕРЖИМОЕ ПРИКРЕПЛЕННОГО ФАЙЛА {doc.file_name}]:\n{extracted_text[-30000:]}"
            else:
                await db_service.log_event(message.from_user.id, "FILE_PARSE_EMPTY", ext)
                await message.answer(i18n.get_text("extract_error", lang))
                return
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
        
        await wait_msg.delete()
        user_query += file_context

    # Режим «роман Толстого»
    if len(user_query) > config.LONG_TEXT_THRESHOLD:
        await db_service.log_event(message.from_user.id, "LONG_INPUT_TRIGGERED")
        await state.update_data(text_buffer=user_query, img_buffer=image_data)
        await state.set_state(BotStates.waiting_for_long_input)
        await message.answer(i18n.get_text("long_input_mode", lang), reply_markup=long_input_kb(lang), parse_mode="Markdown")
        return

    # === [NEW] БУФЕРИЗАЦИЯ И DEBOUNCE ===
    user_id = message.from_user.id
    
    current_text = user_message_buffers.get(user_id, "")
    if current_text:
        user_message_buffers[user_id] = f"{current_text}\n{user_query}".strip()
    else:
        user_message_buffers[user_id] = user_query.strip()
        
    if image_data:
        user_image_buffers[user_id] = image_data

    if user_id in debounce_tasks and not debounce_tasks[user_id].done():
        debounce_tasks[user_id].cancel()
        
    debounce_tasks[user_id] = asyncio.create_task(trigger_processing(message, state, user_id))

@dp.message(BotStates.waiting_for_long_input)
async def long_input_handler(message: types.Message, state: FSMContext):
    data = await state.get_data()
    lang = i18n.user_lang(message.from_user)
    
    if message.text and message.text.strip().lower() in DONE_WORDS:
        await db_service.log_event(message.from_user.id, "LONG_INPUT_FINISHED")
        full_text = data.get("text_buffer", "")
        img_data = data.get("img_buffer", None)
        
        await message.answer(i18n.get_text("doc_accepted", lang), reply_markup=main_kb(lang))
        await state.update_data(text_buffer="", img_buffer=None)
        
        await state.set_state(BotStates.generating)
        try:
            await process_ai_response(message, state, full_text, img_data)
        finally:
            current_state = await state.get_state()
            if current_state == BotStates.generating.state:
                await state.set_state(BotStates.chatting)
        return
    elif message.text in i18n.variants("btn_cancel"):
        await db_service.log_event(message.from_user.id, "LONG_INPUT_CANCELLED")
        await state.update_data(text_buffer="", img_buffer=None)
        await state.set_state(BotStates.chatting)
        await message.answer(i18n.get_text("cancelled", lang), reply_markup=main_kb(lang))
        return

    await db_service.log_event(message.from_user.id, "LONG_INPUT_APPEND")
    current_buffer = data.get("text_buffer", "")
    new_buffer = current_buffer + "\n" + (message.text or "")
    await state.update_data(text_buffer=new_buffer)


async def process_ai_response(message: types.Message, state: FSMContext, user_query: str, image_data: dict = None):
    lang = i18n.user_lang(message.from_user)
    user = await db_service.get_user(message.from_user.id)
    if not user:
        user = await db_service.create_user(message.from_user.id, message.from_user.username, message.from_user.full_name)

    is_unlim = await db_service.has_unlimited(message.from_user.id)
    
    if not is_unlim:
        expired = await db_service.reset_balance_if_expired(message.from_user.id)
        
        if user.credits < 1:
            await db_service.log_event(message.from_user.id, "OUT_OF_CREDITS")
            msg = i18n.get_text("no_credits", lang)
            if expired:
                msg = i18n.get_text("subscription_expired", lang)
            await message.answer(msg)
            return

    wait_msg = await message.answer(i18n.get_text("wait_thinking", lang))

    logger.info(f"📥 [USER {message.from_user.id}]: {user_query[:500]}...")
    await db_service.log_event(message.from_user.id, "USER_MESSAGE", user_query[:4000])

    # === [STATE TRACKER: АФФЕКТИВНАЯ ПАМЯТЬ] ===
    new_states = await db_service.update_user_state(message.from_user.id, fatigue_delta=5)
    
    logger.info(f"📊 [АФФЕКТИВНОЕ ЯДРО] Юзер {message.from_user.id} | AFFECTIVE_STATE_UPDATE trust:{new_states['trust']}, fatigue:{new_states['fatigue']}, mood:{new_states['mood']}")
    
    await db_service.log_event(
        message.from_user.id, 
        "AFFECTIVE_STATE_UPDATE", 
        f"trust:{new_states['trust']}, fatigue:{new_states['fatigue']}, mood:{new_states['mood']}"
    )

    data = await state.get_data()
    history = data.get("history", [])
    
    interaction_count = data.get("interaction_count", 0) + 1

    # === SMART ROUTING ===
    router_decision = await llm_service.analyze_search_need(user_query, history[-4:])
    search_results = ""
    
    if router_decision.get("needs_search"):
        search_query = router_decision.get("query", user_query)
        logger.info(f"🔎 [Router] Решил гуглить: '{search_query}'")
        
        search_results, search_tokens = await llm_service.execute_grounding_search(search_query)
        await db_service.log_event(
            message.from_user.id, 
            "GROUNDING_SEARCH", 
            json.dumps({"query": search_query, "tokens": search_tokens}, ensure_ascii=False)
        )
        logger.info(f"🌐 [Grounding] Факты найдены ({len(search_results)} символов, {search_tokens} токенов)")
    else:
        logger.info("🔎 [Router] Поиск отключен (болтовня — экономим токены)")
    
    # === МНОГОУРОВНЕВАЯ ПАМЯТЬ ===
    rag_ctx, episodic_ctx, semantic_ctx, brief = await asyncio.gather(
        rag_service.search(user_query),
        rag_service.search_episodic_memory(message.from_user.id, user_query),
        semantic_memory.get_facts(message.from_user.id),
        db_service.get_brief(message.from_user.id),
    )
    brief_text = format_brief(brief)
    if brief_text:
        semantic_ctx = brief_text + ("\n\n" + semantic_ctx if semantic_ctx else "")

    # === ЧИТАЕМ УТРЕННЮЮ ПОВЕСТКУ ===
    agenda = {}
    if os.path.exists(AGENDA_PATH):
        try:
            with open(AGENDA_PATH, 'r', encoding='utf-8') as f:
                agenda = json.load(f)
            logger.debug(f"📋 [Повестка] Загружена тема дня: {agenda.get('focus_topic', 'N/A')}")
        except Exception:
            pass

    # === ВНУТРЕННЯЯ ШКАЛА ВРЕМЕНИ ===
    current_time = time.time()
    last_msg_time = data.get("last_message_time")
    
    if not last_msg_time or interaction_count == 1:
        time_passed_str = "Это начало нового разговора."
        delta_seconds = 0
    else:
        delta_seconds = current_time - last_msg_time
        if delta_seconds < 60:
            time_passed_str = "Прошло меньше минуты. Юзер строчит без остановки."
        elif delta_seconds < 3600:
            time_passed_str = f"Прошло {int(delta_seconds // 60)} минут(ы) с прошлого сообщения."
        elif delta_seconds < 86400:
            time_passed_str = f"Прошло {int(delta_seconds // 3600)} час(ов). Юзер вернулся после перерыва."
        else:
            days = int(delta_seconds // 86400)
            if days == 1:
                time_passed_str = "Прошёл 1 день. Юзер пропадал сутки."
            elif days < 7:
                time_passed_str = f"Прошло {days} дня/дней! Юзер где-то пропадал."
            elif days < 30:
                time_passed_str = f"Прошло {days} дней (несколько недель)! Юзер надолго исчез."
            else:
                time_passed_str = f"Прошло {days} дней (МЕСЯЦЫ)! Юзер пропал на огромный срок."

    # ╔══════════════════════════════════════════════════════════════╗
    # ║  [NEW] УМНАЯ АДАПТАЦИЯ ЛОКАЦИИ (Логика в коде, не в промпте)
    # ╚══════════════════════════════════════════════════════════════╝
    base_spatial_context = ContextSimulator.get_current_context()
    last_spatial_context = data.get("last_spatial_context", None)
    
    should_mention_location = False
    
    if interaction_count == 1:
        # Первое сообщение в диалоге — упоминаем локацию
        should_mention_location = True
    elif delta_seconds > 1800:
        # Прошло больше 30 минут — упоминаем локацию
        should_mention_location = True
    elif last_spatial_context and base_spatial_context and last_spatial_context != base_spatial_context:
        # Локация сменилась с прошлого сообщения — упоминаем
        should_mention_location = True
    
    if should_mention_location:
        spatial_context = (
            f"{base_spatial_context}\n"
            "(Ты можешь коротко упомянуть, чем занят, в начале ответа.)"
        )
        logger.debug(f"🏙️ [Среда] Упоминание РАЗРЕШЕНО: {base_spatial_context[:80]}...")
    else:
        spatial_context = (
            f"{base_spatial_context}\n"
            "(Это твой ВНУТРЕННИЙ контекст. НЕ упоминай его в ответе. Отвечай по сути.)"
        )
        logger.debug(f"🏙️ [Среда] Упоминание ЗАПРЕЩЕНО: {base_spatial_context[:80]}...")

    # ╔══════════════════════════════════════════════════════════════╗
    # ║  ШАГ 1: ВНУТРЕННИЙ МОНОЛОГ (THINKING)
    # ╚══════════════════════════════════════════════════════════════╝
    monologue_prompt = Prompts.get_monologue_prompt(
        temp=user.temperature, 
        trust=new_states["trust"], 
        fatigue=new_states["fatigue"], 
        mood=new_states["mood"],
        time_passed=time_passed_str,
        spatial_context=spatial_context
    )
    
    inner_monologue = await llm_service.generate_monologue(
        monologue_prompt, user_query, 
        episodic_context=episodic_ctx, 
        semantic_context=semantic_ctx, 
        history_context=history[-40:], 
        search_results=search_results
    )
    
    await db_service.log_event(message.from_user.id, "INNER_MONOLOGUE", inner_monologue[:4000])
    logger.info(f"🧠 [МОНОЛОГ] Юзер {message.from_user.id} | {time_passed_str} | {inner_monologue[:200]}...")

    # ╔══════════════════════════════════════════════════════════════╗
    # ║  ШАГ 2: ГЕНЕРАЦИЯ ОТВЕТА С ОПОРОЙ НА МОНОЛОГ (ВИДИМЫЙ)
    # ╚══════════════════════════════════════════════════════════════╝
    
    core_beliefs = evolution_engine.get_current_beliefs()
    
    sys_prompt = Prompts.get_system_prompt(
        temp=user.temperature, 
        trust=new_states["trust"], 
        fatigue=new_states["fatigue"], 
        mood=new_states["mood"],
        monologue=inner_monologue,
        core_beliefs=core_beliefs,
        agenda=agenda,
        spatial_context=spatial_context,
        search_results=search_results
    )
    
    failed = False  # при сбое генерации запрос не списываем
    try:
        response_text = await llm_service.generate(
            sys_prompt, user_query, rag_ctx, episodic_ctx, semantic_ctx,
            history[-40:], user.temperature, image_data, search_results
        )
        failed = response_text.startswith(("⚠️", "⛔️"))

        # [NEW] ПЕРЕХВАТЧИК ИГНОРА
        if "[IGNORE]" in response_text:
            logger.info(f"🔇 [ИГНОР] Аристарх молча проигнорировал юзера {message.from_user.id}")
            await db_service.log_event(message.from_user.id, "AI_IGNORED_USER")
            
            new_history = history + [
                {"role": "user", "parts": [{"text": user_query[:4000]}]},
                {"role": "model", "parts": [{"text": "[АРИСТАРХ МОЛЧА ПРОИГНОРИРОВАЛ СООБЩЕНИЕ]"}]}
            ]
            
            await state.update_data(
                history=new_history[-40:], 
                interaction_count=interaction_count,
                last_message_time=current_time,
                last_spatial_context=base_spatial_context  # [NEW] Сохраняем локацию
            )
            
            try:
                await wait_msg.delete()
            except:
                pass
                
            await state.set_state(BotStates.chatting)
            return
        
        if not failed:
            response_text = await polish_reply(llm_service, response_text, user_query)
        response_text = humanize_punctuation(response_text.replace('*', ''))

        logger.info(f"📤 [ARISTARKH {message.from_user.id}]: {response_text[:500]}...")
        
        await db_service.log_event(message.from_user.id, "AI_RESPONSE_TEXT", response_text[:4000])
        await db_service.log_event(message.from_user.id, "AI_RESPONSE_GENERATED", f"chars:{len(response_text)}")
    except Exception as e:
        logger.error(f"Gen Error: {e}")
        response_text = i18n.get_text("generation_error", lang)
        failed = True
        await db_service.log_event(message.from_user.id, "AI_RESPONSE_ERROR", str(e))

    try:
        await wait_msg.delete()
    except:
        pass

    new_history = history + [
        {"role": "user", "parts": [{"text": user_query[:4000]}]},
        {"role": "model", "parts": [{"text": response_text}]}
    ]
    
    # [NEW] Сохраняем last_spatial_context для следующего сообщения
    await state.update_data(
        history=new_history[-40:], 
        interaction_count=interaction_count,
        last_message_time=current_time,
        last_spatial_context=base_spatial_context
    )
    
    if interaction_count % 5 == 0:
        asyncio.create_task(extract_and_save_memory(message.from_user.id, new_history[-10:]))
        await db_service.log_event(
            message.from_user.id, 
            "MEMORY_CONSOLIDATION_TRIGGERED", 
            f"interaction:{interaction_count}, history_len:{len(new_history)}"
        )
    
    if not is_unlim and not failed:
        await db_service.update_balance(message.from_user.id, -1)

    await send_smart_response(message, response_text)

    # Один раз за сессию напоминаем про бриф тем, кто его ещё не заполнил
    if not brief and not failed and not data.get("brief_hint_shown"):
        await message.answer(i18n.get_text("brief_hint", lang))
        await state.update_data(brief_hint_shown=True)

    try:
        suggestion = await assistant_service.generate_suggestion(new_history[-40:])
        if suggestion:
            clean_suggestion = html.escape(suggestion.strip())
            suggestion_msg = i18n.get_text("assistant_hint", lang, text=clean_suggestion)
            await message.answer(suggestion_msg, parse_mode="HTML")
            await db_service.log_event(message.from_user.id, "ASSISTANT_SUGGESTION_SENT", suggestion)
    except Exception as e:
        logger.error(f"Assistant generation failed: {e}")
        await db_service.log_event(message.from_user.id, "ASSISTANT_ERROR", str(e))


async def extract_and_save_memory(user_id: int, history_chunk: list):
    try:
        mem_data = await assistant_service.extract_memory_and_state(history_chunk)
        if not mem_data: 
            logger.warning(f"⚠️ Экстрактор вернул пустой результат для юзера {user_id}")
            return
        
        if mem_data.get("summary"):
            await rag_service.add_episodic_memory(
                user_id, mem_data["summary"], mem_data.get("emotion_weight", 0.5)
            )
            logger.info(f"🧠 [Эпизодическая] Сохранено саммари для {user_id}")
            
        if mem_data.get("facts"):
            await semantic_memory.add_facts(user_id, mem_data["facts"])
            logger.info(f"🕸️ [Семантическая] Сохранено {len(mem_data['facts'])} фактов для {user_id}")
            
        if mem_data.get("trust_delta") or mem_data.get("mood"):
            await db_service.update_user_state(
                user_id, 
                trust_delta=mem_data.get("trust_delta", 0), 
                new_mood=mem_data.get("mood")
            )
            logger.info(f"💓 [Аффективная] Обновлены trust={mem_data.get('trust_delta', 0)}, mood={mem_data.get('mood')} для {user_id}")
            
        await db_service.log_event(
            user_id, "MEMORY_CONSOLIDATED", 
            f"facts:{len(mem_data.get('facts', []))}, summary:{bool(mem_data.get('summary'))}, "
            f"trust_delta:{mem_data.get('trust_delta', 0)}, mood:{mem_data.get('mood', 'N/A')}"
        )
        logger.info(f"🔄 Память полностью консолидирована для юзера {user_id}")
        
    except Exception as e:
        logger.error(f"❌ Ошибка фонового воркера памяти для {user_id}: {e}")
        await db_service.log_event(user_id, "MEMORY_CONSOLIDATION_ERROR", str(e))


async def main():
    logger.info("🚀 Запуск Аристарха v22.0 (Умная локация + Защита персоны)...")
    await db_service.init_db()
    await rag_service.build_index()
    
    http_session = aiohttp.ClientSession()
    llm_service.http_session = http_session
    rag_service.http_session = http_session
    assistant_service.http_session = http_session 
    evolution_engine.http_session = http_session
    
    try:
        try:
            await bot.delete_webhook(drop_pending_updates=True)
        except Exception as e:
            logger.error(f"Не удалось сбросить вебхук из-за таймаута Telegram. Идем дальше: {e}")
            
        await dp.start_polling(bot)
    finally:
        await http_session.close()

if __name__ == "__main__":
    try: asyncio.run(main())
    except KeyboardInterrupt: pass